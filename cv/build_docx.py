"""Bangun CV_Steven_Huang_ATS.docx dari cv.html (satu sumber isi).

Pakai: python build_docx.py   (butuh: pip install python-docx beautifulsoup4)
"""
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

HERE = Path(__file__).parent
FONT = "Arial"
RIGHT_TAB = Mm(180)  # lebar A4 (210) dikurangi margin kiri-kanan 15 mm


def text(el):
    return " ".join(el.get_text(" ", strip=True).split())


def base_paragraph(doc, size=10, space_after=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    return p


def run(p, s, size=10, bold=False, italic=False):
    r = p.add_run(s)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return r


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def inline_runs(p, el, size=10):
    """Tulis isi elemen, mempertahankan <strong> sebagai bold."""
    for node in el.children:
        if isinstance(node, NavigableString):
            s = " ".join(str(node).split())
            if s:
                lead = " " if str(node)[:1].isspace() and p.runs else ""
                run(p, lead + s, size)
        else:
            s = text(node)
            lead = " " if p.runs and not p.runs[-1].text.endswith(" ") else ""
            run(p, lead + s, size, bold=node.name == "strong")


def main():
    soup = BeautifulSoup((HERE / "cv.html").read_text(encoding="utf-8"), "html.parser")
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    sec.top_margin, sec.bottom_margin = Mm(12), Mm(11)
    sec.left_margin = sec.right_margin = Mm(15)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)

    for el in soup.body.find_all(recursive=False):
        cls = el.get("class") or []
        if el.name == "h1":
            run(base_paragraph(doc, space_after=2), text(el), 20, bold=True)
        elif el.name == "p" and "contact" in cls:
            run(base_paragraph(doc), text(el), 9.5)
        elif el.name == "h2":
            p = base_paragraph(doc, space_after=3)
            p.paragraph_format.space_before = Pt(8)
            run(p, text(el).upper(), 10.5, bold=True)
            bottom_border(p)
        elif el.name == "p":
            run(base_paragraph(doc, space_after=2), text(el))
        elif el.name == "div" and "entry" in cls:
            for child in el.find_all(recursive=False):
                if child.name == "div" and "row" in (child.get("class") or []):
                    left, *right = child.find_all(recursive=False)
                    p = base_paragraph(doc, space_after=0)
                    run(p, text(left), 10.2, bold=left.name == "strong", italic="sub" in (left.get("class") or []))
                    if right:
                        run(p, "\t" + text(right[0]))
                elif child.name == "div":
                    run(base_paragraph(doc, space_after=0), text(child), italic=True)
                elif child.name == "ul":
                    add_list(doc, child)
        elif el.name == "ul":
            add_list(doc, el)
        elif el.name == "div" and "skills" in cls:
            for p_el in el.find_all("p"):
                inline_runs(base_paragraph(doc, space_after=1), p_el)

    out = HERE / "CV_Steven_Huang_ATS.docx"
    doc.save(out)
    print("Tersimpan:", out)


def add_list(doc, ul):
    for li in ul.find_all("li", recursive=False):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.08
        inline_runs(p, li)


if __name__ == "__main__":
    main()
